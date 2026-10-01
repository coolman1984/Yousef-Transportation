"""Excel in and out: reading the workbook, the preview with its messages, the confirmed import, the round trip to the
"same as today" layout, the clean export, and every refusal."""
import datetime as dt
import io
import os
import sys
import unittest

from harness import ApiError, Server, make_authority

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'server'))
import domain  # noqa: E402
import excel_io as X  # noqa: E402
import xlsx_read  # noqa: E402
import xlsx_write as w  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLE = os.path.join(HERE, 'fixtures', 'sample_sep26_synthetic.xlsx')


def sample():
    with open(SAMPLE, 'rb') as f:
        return f.read()


class ReadAndPlanTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows, cls.report = X.read_workbook(sample())
        cls.plan = X.plan(cls.rows, {})

    def test_sheets_and_rows(self):
        self.assertEqual([(r['sheet'], r['rows']) for r in self.report], [('All Car', 232), ('SUV Rent', 11), ('Microbus Rent', 4)])
        self.assertEqual(len(self.rows), 247)          # totals rows and blank template rows are not trips

    def test_names_are_cleaned_and_counted_once(self):
        self.assertEqual(len(self.plan['new']['drivers']), 41)     # "  Ali Hassan " and "Ali Hassan" are one driver
        self.assertEqual(len({d['plate'] for d in self.plan['new']['vehicles']}), len(self.plan['new']['vehicles']))
        leading = [r for r in self.rows if r['plate'].startswith(' ')]
        self.assertEqual(leading, [])                                # the leading space was cleaned when reading

    def test_problem_rows_and_messages(self):
        st = self.plan['stats']
        self.assertEqual((st['total'], st['new'] + st['problem'], st['duplicate']), (247, 247, 0))
        self.assertEqual(st['problem'], 7)                          # the rows without a plate
        codes = {m['code'] for it in self.plan['rows'] for m in it['messages']}
        self.assertIn('plate', codes)
        self.assertIn('km', codes)

    def test_alerts(self):
        alerts = {a['code']: a for a in self.plan['alerts']}
        self.assertGreaterEqual(alerts['backsteps']['n'], 8)         # the odometer faults planted in the sample
        self.assertEqual(alerts['backsteps']['level'], 'bad')
        self.assertIn('no_times', alerts)
        self.assertIn('billable', alerts)

    def test_the_31000_km_typo_is_found(self):
        big = [it for it in self.plan['rows'] if any(m['code'] == 'backstep' and '3' in m['text'] and int(m['text'].split()[3]) > 20000 for m in it['messages'])]
        self.assertTrue(big)

    def test_spellings_of_a_place_are_suggested_for_merging(self):
        flat = [v.lower() for m in self.plan['merges'] for v in [m['canonical']] + m['variants']]
        self.assertTrue(any('musiem' in v or 'museum' in v for v in flat), self.plan['merges'])

    def test_rent_sheets_take_the_category_from_the_sheet(self):
        cats = {c['name'] for c in self.plan['new']['categories']}
        self.assertTrue({'Manager car', 'Extra', 'SUV Rent', 'Microbus Rent'} <= cats)


class RefusalTest(unittest.TestCase):
    def test_messages(self):
        for data, part in ((b'', 'empty'), (b'\x00\x01\x02' * 300, 'not recognised'), (w.build([w.Sheet('A', [['a', 'b'], [1, 2]])]), 'no sheet with trip columns'),
                           (w.build([w.Sheet('A', [['Date', 'Driver Name', 'Car Plate']])]), 'no trips')):
            with self.assertRaises(X.ImportError_) as e:
                X.read_workbook(data)
            self.assertIn(part, str(e.exception).lower())

    def test_bad_cells_become_row_messages_not_crashes(self):
        data = w.build([w.Sheet('A', [['Date', 'Driver Name', 'Car Plate', 'Strat KM', 'End KM'], ['not a date', 'X Y', 'ط و ي 6829', 'abc', 5], [dt.date(2026, 9, 1), 'X Y', 'ط و ي 6829', 200, 100]])])
        rows, _ = X.read_workbook(data)
        p = X.plan(rows, {})
        self.assertEqual([it['status'] for it in p['rows']], ['problem', 'new'])
        self.assertIn('date', {m['code'] for m in p['rows'][0]['messages']})
        self.assertIn('km_order', {m['code'] for m in p['rows'][1]['messages']})


class ApiImportExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.S = Server('excel').start()
        cls.ac = make_authority(cls.S)

    @classmethod
    def tearDownClass(cls):
        cls.S.cleanup()

    def raw(self, c, path, data):
        return c.call('POST', path, raw=data, headers={'Content-Type': 'application/octet-stream'})

    def test_a_flow(self):
        ac = self.ac
        pv = self.raw(ac, '/api/excel/preview?name=sep.xlsx', sample())
        self.assertEqual(pv['stats']['total'], 247)
        self.assertEqual(ac.get('/api/state')['trips'], [], 'the preview writes nothing')
        res = ac.post('/api/excel/commit', {'id': pv['id'], 'includeProblems': True,
                                            'merges': [m for m in pv['merges']]})
        self.assertEqual(res['trips'], 247)
        st = ac.get('/api/state')
        self.assertEqual(len(st['trips']), 247)
        self.assertTrue(all(t['source'] == 'excel' and t['locked'] and t['status'] == 'closed' for t in st['trips']))
        self.assertEqual(len(st['drivers']), 41)
        self.assertEqual(len({t['no'] for t in st['trips']}), 247, 'every trip has its own number')
        self.assertEqual(sorted(c['exportSheet'] for c in st['tripCategories'] if c['name'] in ('SUV Rent', 'Microbus Rent')), ['Microbus Rent', 'SUV Rent'])
        audit = ac.get('/api/audit?limit=5')['rows']
        self.assertTrue(any('Excel import' in r['label'] for r in audit))
        # the same file again: everything is a duplicate, nothing is written
        pv2 = self.raw(ac, '/api/excel/preview?name=sep.xlsx', sample())
        self.assertEqual((pv2['stats']['duplicate'], pv2['stats']['new']), (247, 0))
        with self.assertRaises(ApiError) as e:
            ac.post('/api/excel/commit', {'id': pv2['id']})
        self.assertEqual(e.exception.code, 400)
        self.assertIn('nothing to import', str(e.exception.msg).lower())
        self.assertEqual(len(ac.get('/api/state')['trips']), 247)

    def test_b_round_trip_to_the_same_layout(self):
        ac = self.ac
        st = ac.get('/api/state')
        if not st['trips']:
            self.skipTest('needs test_a')
        data = ac.call('GET', '/api/excel/export?ym=2026-09&layout=today')
        wb = xlsx_read.read(data)
        self.assertEqual([s.name for s in wb.sheets], ['All Car', 'SUV Rent', 'Microbus Rent'])
        allc = wb.sheet('All Car')
        self.assertEqual([allc.value(1, c) for c in range(1, 17)], X.TODAY_HEADERS)
        self.assertEqual(allc.tables[0]['name'], 'Table1')
        self.assertEqual(allc.tables[0]['columns'][11], ' KM')
        original = X.read_workbook(sample())[0]
        want = sorted((str(r['date']), r['driver'], domain.norm_plate(r['plate'])[0] if r['plate'] else '', r['requester'], r['category'], r['destination'],
                       r['startKm'], r['endKm'], r['billableKm']) for r in original if r['sheet'] == 'All Car')
        got = []
        for r in range(2, allc.max_row + 1):
            v = lambda c: allc.value(r, c)  # noqa: E731
            got.append((str(v(1)), v(2), v(3) or '', v(4) or '', v(6), v(8), v(10), v(11), v(13)))
        self.assertEqual(sorted(got), want)
        # formulas are there and the cached values are right
        r = 2
        self.assertTrue(allc.cells[(r, 12)].f.startswith('+Table1[[#This Row],[End KM]]'))
        if allc.value(r, 10) is not None:
            self.assertEqual(allc.value(r, 12), allc.value(r, 11) - allc.value(r, 10))
        self.assertTrue(allc.cells[(r, 16)].f.startswith('IF(O2-N2>TIME(12,0,0)'))
        suv = wb.sheet('SUV Rent')
        self.assertEqual(sum(1 for r in range(2, suv.max_row + 1) if suv.value(r, 1)), 11)
        self.assertEqual(suv.value(suv.max_row, 9), 'TTL')
        self.assertTrue(suv.cells[(suv.max_row, 11)].f.startswith('+SUM(K2:K12)'))
        bus = wb.sheet('Microbus Rent')
        self.assertEqual(bus.tables[0]['name'], 'Table2')
        self.assertEqual(bus.value(bus.max_row, 10), 'TTL')
        self.assertTrue(bus.cells[(bus.max_row, 12)].f.startswith('SUBTOTAL(109,Table2'))
        try:
            import openpyxl
        except ImportError:
            return
        book = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(sorted(book['All Car'].tables), ['Table1'])

    def test_c_clean_export_and_template(self):
        for lang in ('en', 'ar'):
            data = self.ac.call('GET', f'/api/excel/export?ym=2026-09&layout=clean&lang={lang}')
            wb = xlsx_read.read(data)
            names = [s.name for s in wb.sheets]
            self.assertEqual(len(names), 10, names)
            trips = wb.sheets[1]
            self.assertEqual(trips.max_row - 1, 247)
            self.assertIsInstance(trips.value(2, 2), dt.date)
        tpl = xlsx_read.read(self.ac.call('GET', '/api/excel/template'))
        self.assertEqual([tpl.sheets[0].value(1, c) for c in range(1, 4)], ['Date', 'Driver Name', 'Car Plate'])

    def test_d_permissions_and_bad_input(self):
        ac = self.ac
        for path, body in (('/api/excel/preview', b'not excel'),):
            with self.assertRaises(ApiError) as e:
                self.raw(ac, path, body)
            self.assertEqual(e.exception.code, 400)
        with self.assertRaises(ApiError) as e:
            ac.get('/api/excel/export?ym=bad&layout=today')
        self.assertEqual(e.exception.code, 400)
        with self.assertRaises(ApiError):
            ac.post('/api/excel/commit', {'id': 'nope'})
        viewer = next(p for p in ac.get('/api/users')['profiles'] if p['id'] == 'viewer')
        ac.post('/api/users/save', {'username': 'vic.v', 'full_name': 'Vic Viewer', 'password': 'Quarter-pass77', 'must_change': False, 'role': viewer['name'], 'perms': viewer['perms'], 'scopes': None})
        c = self.S.client()
        c.login('vic.v', 'Quarter-pass77')
        for fn in (lambda: self.raw(c, '/api/excel/preview', sample()), lambda: c.get('/api/excel/template')):
            with self.assertRaises(ApiError) as e:
                fn()
            self.assertEqual(e.exception.code, 403)


if __name__ == '__main__':
    unittest.main()
