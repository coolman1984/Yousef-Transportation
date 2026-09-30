"""The dependency-free Excel reader and writer: types, formulas, tables, dates, damaged files, and a real-file cross-check."""
import datetime as dt
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
import io

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'server'))
import xlsx_read  # noqa: E402
import xlsx_write as w  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REAL = os.path.join(ROOT, 'samples', 'private', 'Extra_Sep-26_Recovered.xlsx')


def sample_sheet():
    head = ['Date', 'Name', 'Strat KM', 'End KM', ' KM', 'Start', 'End', 'OT Hours']
    rows = [head]
    for i, (s, e, a, b) in enumerate([(100, 150, dt.time(6, 30), dt.time(19, 30)), (150, 300, dt.time(7, 0), dt.time(20, 0))], start=2):
        rows.append([dt.date(2026, 9, i), 'ط و ي 6829 - Ali', s, e,
                     w.Formula('+Table1[[#This Row],[End KM]]-Table1[[#This Row],[Strat KM]]', e - s), a, b,
                     w.Formula(f'IF(G{i}-F{i}>TIME(12,0,0),G{i}-F{i}-TIME(12,0,0),0)', dt.time(1, 0) and 0.0416666667)])
    rows.append([None, None, None, w.S('TTL', 'total_text'), w.S(w.Formula('SUBTOTAL(109,Table1[[ KM]])', 200), 'total_int')])
    t = w.Table('Table1', 'A1:H3', head, formulas={' KM': '+Table1[[#This Row],[End KM]]-Table1[[#This Row],[Strat KM]]'})
    return w.Sheet('All Car', rows, widths=[12, 20, 10, 10, 8, 8, 8, 10], tables=[t], freeze=(1, 0), grid=False)


class RoundTripTest(unittest.TestCase):
    def setUp(self):
        self.data = w.build([sample_sheet(), w.Sheet('Second', [['a', 1], ['b', 2.5], [True, dt.datetime(2026, 9, 1, 8, 15)]])])
        self.wb = xlsx_read.read(self.data)

    def test_values_types_and_formulas_survive(self):
        s = self.wb.sheet('All Car')
        self.assertEqual(s.value(2, 1), dt.date(2026, 9, 2))
        self.assertEqual(s.value(2, 2), 'ط و ي 6829 - Ali')
        self.assertEqual((s.value(2, 3), s.value(2, 4), s.value(2, 5)), (100, 150, 50))
        self.assertEqual(s.value(2, 6), dt.time(6, 30))
        self.assertTrue(s.cells[(2, 5)].f.startswith('+Table1[[#This Row]'))
        self.assertEqual(s.value(4, 5), 200)
        s2 = self.wb.sheet('Second')
        self.assertEqual((s2.value(1, 1), s2.value(1, 2), s2.value(2, 2), s2.value(3, 1), s2.value(3, 2)), ('a', 1, 2.5, True, dt.datetime(2026, 9, 1, 8, 15)))

    def test_table_and_layout_parts(self):
        t = self.wb.sheet('All Car').tables[0]
        self.assertEqual((t['name'], t['ref']), ('Table1', 'A1:H3'))
        self.assertEqual(t['columns'][4], ' KM')   # the leading space of the original header is kept
        z = zipfile.ZipFile(io.BytesIO(self.data))
        xml = z.read('xl/worksheets/sheet1.xml').decode()
        self.assertIn('showGridLines="0"', xml)
        self.assertIn('<pane xSplit="1"', xml)
        self.assertIn('calculatedColumnFormula', z.read('xl/tables/table1.xml').decode())
        self.assertIn('fullCalcOnLoad="1"', z.read('xl/workbook.xml').decode())

    def test_openpyxl_accepts_our_file(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl not installed')
        wb = openpyxl.load_workbook(io.BytesIO(self.data))
        ws = wb['All Car']
        self.assertEqual(list(ws.tables), ['Table1'])
        self.assertEqual(ws.tables['Table1'].ref, 'A1:H3')
        self.assertEqual(ws['E2'].value, '=+Table1[[#This Row],[End KM]]-Table1[[#This Row],[Strat KM]]')
        self.assertEqual(ws.freeze_panes, 'B1')
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertEqual(ws['A2'].number_format, 'mm-dd-yy')
        self.assertEqual(ws['F2'].number_format, 'h:mm')

    def test_special_characters_are_escaped(self):
        data = w.build([w.Sheet('A/B:C?', [['<b>&"x"', 'line\x01break', 'ok']])])
        wb = xlsx_read.read(data)
        self.assertEqual(wb.sheets[0].name, 'A B C')
        self.assertEqual(wb.sheets[0].value(1, 1), '<b>&"x"')
        self.assertEqual(wb.sheets[0].value(1, 2), 'linebreak')

    def test_duplicate_sheet_names_are_made_unique(self):
        wb = xlsx_read.read(w.build([w.Sheet('X', [[1]]), w.Sheet('x', [[2]])]))
        self.assertEqual([s.name for s in wb.sheets], ['X', 'x 2'])


class ReaderRefusalTest(unittest.TestCase):
    def test_clear_messages(self):
        for data, part in ((b'', 'empty'), (b'not a workbook at all' * 20, 'not an excel'), (b'PK' + b'\x00' * 200, 'damaged')):
            with self.assertRaises(xlsx_read.XlsxError) as e:
                xlsx_read.read(data)
            self.assertIn(part, str(e.exception).lower())

    def test_zip_without_workbook(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z:
            z.writestr('hello.txt', 'x' * 200)
        with self.assertRaises(xlsx_read.XlsxError):
            xlsx_read.read(buf.getvalue())

    def test_date_serials(self):
        self.assertEqual(xlsx_read.excel_date(61), dt.datetime(1900, 3, 1))     # the fake 29 Feb 1900 is skipped
        self.assertEqual(xlsx_read.excel_date(46023), dt.datetime(2026, 1, 1))
        self.assertEqual(xlsx_read.excel_date(0.5), dt.datetime(1899, 12, 31, 12))
        self.assertEqual(xlsx_read.excel_date(1, date1904=True), dt.datetime(1904, 1, 2))


def libreoffice_reads_xlsx():
    """True when LibreOffice here can load an .xlsx at all (some minimal installs cannot); then it is a second opinion on our files."""
    if not shutil.which('soffice'):
        return False
    try:
        import openpyxl
    except ImportError:
        return False
    d = tempfile.mkdtemp()
    try:
        wb = openpyxl.Workbook()
        wb.active.append(['a', 1])
        wb.save(os.path.join(d, 'ref.xlsx'))
        subprocess.run(['soffice', '--headless', '--convert-to', 'csv', '--outdir', d, os.path.join(d, 'ref.xlsx')], capture_output=True, timeout=120, env=dict(os.environ, HOME=d))
        return os.path.exists(os.path.join(d, 'ref.csv'))
    except Exception:  # noqa: BLE001
        return False
    finally:
        shutil.rmtree(d, ignore_errors=True)


class OpensInLibreOfficeTest(unittest.TestCase):
    def test_file_opens_and_recalculates(self):
        if not libreoffice_reads_xlsx():
            self.skipTest('LibreOffice cannot read .xlsx on this machine')
        d = tempfile.mkdtemp()
        try:
            p = os.path.join(d, 'a.xlsx')
            with open(p, 'wb') as f:
                f.write(w.build([sample_sheet()]))
            subprocess.run(['soffice', '--headless', '--convert-to', 'csv', '--outdir', d, p], capture_output=True, timeout=120, env=dict(os.environ, HOME=d))
            with open(os.path.join(d, 'a.csv'), encoding='utf-8') as f:
                csv = f.read().splitlines()
            self.assertEqual(csv[0].split(',')[:3], ['Date', 'Name', 'Strat KM'])
            self.assertIn('50', csv[1].split(','))          # the table formula was recalculated
            self.assertIn('150', csv[2].split(','))
        finally:
            shutil.rmtree(d, ignore_errors=True)


@unittest.skipUnless(os.path.exists(REAL), 'the owner\'s workbook is not in samples/private')
class RealFileTest(unittest.TestCase):
    def test_matches_openpyxl_cell_by_cell(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl not installed')
        wb = xlsx_read.read(open(REAL, 'rb').read())
        ref = openpyxl.load_workbook(REAL, data_only=True)
        diffs = 0
        for s in wb.sheets:
            for (r, c), cell in s.cells.items():
                ov, v = ref[s.name].cell(r, c).value, cell.v
                same = ov == v or (type(v) is dt.date and isinstance(ov, dt.datetime) and ov.date() == v) or \
                    (isinstance(ov, (int, float)) and isinstance(v, (int, float)) and abs(ov - v) < 1e-9)
                diffs += not same
        self.assertEqual(diffs, 0)
        self.assertEqual([s.name for s in wb.sheets], ['All Car', 'SUV Rent', 'Microbus Rent'])
        self.assertEqual(wb.sheet('All Car').tables[0]['name'], 'Table1')


if __name__ == '__main__':
    unittest.main()
